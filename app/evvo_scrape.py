import re

import pandas as pd

from parsing import safe_strip


def evvo_scrape():
    import re
    import time
    from urllib.parse import urljoin

    import pandas as pd
    import requests
    from bs4 import BeautifulSoup
    from openpyxl import load_workbook
    from openpyxl.styles import Font


    TVS_URL = (
        "https://www.evvosemi.com/"
        "en/index/products/pid/70/ty/77.html"
    )

    ZENER_URL = (
        "https://www.evvosemi.com/"
        "en/index/products/pid/70/ty/106/iscat/1.html"
    )

    TVS_OUTPUT = "evvo_tvs_specs.xlsx"
    ZENER_OUTPUT = "evvo_zener_specs.xlsx"

    REQUEST_DELAY = 0.20
    TIMEOUT = 30

    HEADERS = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/151.0.0.0 Safari/537.36"
        ),
        "Accept": (
            "text/html,application/xhtml+xml,"
            "application/xml;q=0.9,"
            "image/avif,image/webp,image/apng,*/*;q=0.8"
        ),
        "Accept-Language": "en-US,en;q=0.9",
        "Upgrade-Insecure-Requests": "1",
    }

    OUTPUT_COLUMNS = [
        "Model Number",
        "Encapsulation",
        "Package",
        "Description",
        "Pin-To-Pin Flat Type",
        "Data Sheet",
    ]


    def clean_text(value):
        if value is None:
            return ""

        return re.sub(
            r"\s+",
            " ",
            str(value),
        ).strip()


    def extract_datasheet_url(td, page_url):
        if td is None:
            return ""

        for a in td.find_all(
            "a",
            href=True,
        ):
            href = clean_text(
                a.get("href")
            )

            if not href:
                continue

            if (
                ".pdf" in href.lower()
                or "download" in href.lower()
                or "datasheet" in href.lower()
            ):
                return urljoin(
                    page_url,
                    href,
                )

        return ""


    def parse_product_row(
        tr,
        page_url,
    ):
        """
        EVVO product table:

          0 Product picture
          1 Model number
          2 encapsulation
          3 package
          4 Description
          5 Pin-To-Pin Flat type
          6 Data sheet
        """

        tds = tr.find_all(
            "td",
            recursive=False,
        )

        if len(tds) < 7:
            tds = tr.find_all(
                "td"
            )

        if len(tds) < 7:
            return None

        model_number = clean_text(
            tds[1].get_text(
                " ",
                strip=True,
            )
        )

        if not model_number:
            return None

        if (
            model_number.lower()
            == "model number"
        ):
            return None

        if not re.search(
            r"[A-Za-z0-9]",
            model_number,
        ):
            return None

        return {
            "Model Number":
                model_number,

            "Encapsulation":
                clean_text(
                    tds[2].get_text(
                        " ",
                        strip=True,
                    )
                ),

            "Package":
                clean_text(
                    tds[3].get_text(
                        " ",
                        strip=True,
                    )
                ),

            "Description":
                clean_text(
                    tds[4].get_text(
                        " ",
                        strip=True,
                    )
                ),

            "Pin-To-Pin Flat Type":
                clean_text(
                    tds[5].get_text(
                        " ",
                        strip=True,
                    )
                ),

            "Data Sheet":
                extract_datasheet_url(
                    tds[6],
                    page_url,
                ),
        }


    def get_page(
        session,
        base_url,
        page_number,
    ):
        url = (
            f"{base_url}"
            f"?page={page_number}"
        )

        headers = dict(
            HEADERS
        )

        headers["Referer"] = (
            base_url
        )

        response = session.get(
            url,
            headers=headers,
            timeout=TIMEOUT,
        )

        response.raise_for_status()

        if not response.encoding:
            response.encoding = (
                "utf-8"
            )

        return (
            url,
            response.text,
        )


    def detect_last_page(
        html,
    ):
        soup = BeautifulSoup(
            html,
            "html.parser",
        )

        page_numbers = [1]

        for a in soup.find_all(
            "a",
            href=True,
        ):
            href = a.get(
                "href",
                "",
            )

            match = re.search(
                r"[?&]page=(\d+)",
                href,
                flags=re.IGNORECASE,
            )

            if match:
                page_numbers.append(
                    int(
                        match.group(1)
                    )
                )

        return max(
            page_numbers
        )


    def scrape_page(
        html,
        page_url,
    ):
        soup = BeautifulSoup(
            html,
            "html.parser",
        )

        records = []

        for tr in soup.find_all(
            "tr"
        ):
            record = (
                parse_product_row(
                    tr,
                    page_url,
                )
            )

            if record:
                records.append(
                    record
                )

        return records


    def save_excel(
        records,
        output_file,
        sheet_name,
    ):
        df = pd.DataFrame(
            records
        )

        for col in OUTPUT_COLUMNS:
            if col not in df.columns:
                df[col] = ""

        df = df[
            OUTPUT_COLUMNS
        ]

        df = (
            df
            .drop_duplicates(
                subset=[
                    "Model Number"
                ],
                keep="first",
            )
            .reset_index(
                drop=True
            )
            .fillna("")
        )

        df.to_excel(
            output_file,
            index=False,
        )

        wb = load_workbook(
            output_file
        )

        ws = wb.active

        ws.title = (
            sheet_name
        )

        ws.freeze_panes = (
            "A2"
        )

        ws.auto_filter.ref = (
            ws.dimensions
        )

        for cell in ws[1]:
            cell.font = Font(
                bold=True
            )

        for column_cells in (
            ws.columns
        ):
            letter = (
                column_cells[0]
                .column_letter
            )

            max_length = 0

            for cell in (
                column_cells
            ):
                value = (
                    ""
                    if cell.value is None
                    else str(
                        cell.value
                    )
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            width = min(
                max(
                    max_length + 2,
                    14,
                ),
                60,
            )

            ws.column_dimensions[
                letter
            ].width = width

        datasheet_col = (
            OUTPUT_COLUMNS.index(
                "Data Sheet"
            )
            + 1
        )

        for row in range(
            2,
            ws.max_row + 1,
        ):
            cell = ws.cell(
                row=row,
                column=datasheet_col,
            )

            if (
                isinstance(
                    cell.value,
                    str,
                )
                and cell.value.startswith(
                    "http"
                )
            ):
                cell.hyperlink = (
                    cell.value
                )

                cell.style = (
                    "Hyperlink"
                )

        wb.save(
            output_file
        )

        return df


    def scrape_catalog(
        session,
        catalog_name,
        base_url,
        output_file,
    ):
        print()
        print("=" * 70)
        print(
            f"EVVO {catalog_name} SCRAPER"
        )
        print("=" * 70)

        print(
            f"URL: {base_url}"
        )

        print(
            f"Output: {output_file}"
        )

        print()

        print(
            "Downloading page 1..."
        )

        (
            page1_url,
            page1_html,
        ) = get_page(
            session,
            base_url,
            1,
        )

        last_page = (
            detect_last_page(
                page1_html
            )
        )

        print(
            f"Detected "
            f"{last_page} pages."
        )

        page1_records = (
            scrape_page(
                page1_html,
                page1_url,
            )
        )

        print(
            f"Page 1/"
            f"{last_page}: "
            f"{len(page1_records)} "
            f"rows found."
        )

        if not page1_records:
            raise RuntimeError(
                f"No EVVO "
                f"{catalog_name} "
                f"products were "
                f"parsed from page 1."
            )

        print()
        print(
            "First parsed row:"
        )

        for key in (
            OUTPUT_COLUMNS
        ):
            print(
                f"  {key}: "
                f"{page1_records[0].get(key, '')}"
            )

        all_records = list(
            page1_records
        )

        seen = {
            row["Model Number"]
            for row
            in page1_records
        }

        for page_number in range(
            2,
            last_page + 1,
        ):
            (
                page_url,
                html,
            ) = get_page(
                session,
                base_url,
                page_number,
            )

            records = (
                scrape_page(
                    html,
                    page_url,
                )
            )

            new_records = []

            for record in records:
                model = (
                    record[
                        "Model Number"
                    ]
                )

                if model not in seen:
                    seen.add(
                        model
                    )

                    new_records.append(
                        record
                    )

            all_records.extend(
                new_records
            )

            print(
                f"Page "
                f"{page_number}/"
                f"{last_page}: "
                f"{len(records)} rows, "
                f"{len(new_records)} new, "
                f"{len(all_records)} total"
            )

            time.sleep(
                REQUEST_DELAY
            )

        print()

        print(
            f"Saving "
            f"{len(all_records)} "
            f"unique products..."
        )

        df = save_excel(
            all_records,
            output_file,
            f"EVVO {catalog_name}",
        )

        print(
            f"Finished "
            f"{catalog_name}."
        )

        print(
            f"Saved "
            f"{len(df)} rows."
        )

        print(
            f"Output: "
            f"{output_file}"
        )

        return df


    session = (
        requests.Session()
    )

    # ========================================================
    # ESD / TVS
    # ========================================================

    tvs_df = scrape_catalog(
        session=session,

        catalog_name=(
            "ESD TVS"
        ),

        base_url=(
            TVS_URL
        ),

        output_file=(
            TVS_OUTPUT
        ),
    )


    # ========================================================
    # ZENER
    # ========================================================

    zener_df = scrape_catalog(
        session=session,

        catalog_name=(
            "Zener"
        ),

        base_url=(
            ZENER_URL
        ),

        output_file=(
            ZENER_OUTPUT
        ),
    )


    print()
    print("=" * 70)
    print("ALL EVVO SCRAPES COMPLETE")
    print("=" * 70)

    print(
        f"{TVS_OUTPUT}: "
        f"{len(tvs_df)} rows"
    )

    print(
        f"{ZENER_OUTPUT}: "
        f"{len(zener_df)} rows"
    )

    return (
        tvs_df,
        zener_df,
    )


# ============================================================
# SPEC COLUMNS
#
# The saved sheets' header names, kept as constants so the
# parsers below read the columns by name instead of by
# position. They match OUTPUT_COLUMNS exactly.
#
# Note that "Package" is EVVO's shipping format (T/R), not a
# body. The body is in "Encapsulation".
# ============================================================

MODEL_NUMBER_COLUMN = "Model Number"

ENCAPSULATION_COLUMN = "Encapsulation"

SHIPPING_PACKAGE_COLUMN = "Package"

DESCRIPTION_COLUMN = "Description"

PIN_TO_PIN_COLUMN = "Pin-To-Pin Flat Type"

DATASHEET_COLUMN = "Data Sheet"


# ============================================================
# SPEC HELPERS
#
# EVVO publishes no parametric columns at all - every number
# lives in the free-text Description, so the helpers below
# pull them back out of it.
# ============================================================

def get_cell(
    part_row,
    column_name,
):

    if column_name in part_row.index:
        return part_row[column_name]

    wanted = str(column_name).strip().lower()

    for column in part_row.index:

        if str(column).strip().lower() == wanted:
            return part_row[column]

    return None


def get_text(
    part_row,
    column_name,
):

    value = get_cell(
        part_row,
        column_name,
    )

    if value is None or (
        not isinstance(value, str)
        and pd.isna(value)
    ):
        return ""

    return safe_strip(value)


def get_power(
    description
):

    # --------------------------------------------------------
    # "600W", "300MW" and "Ptot= 350mW" all appear. Returned
    # in watts.
    # --------------------------------------------------------

    power_match = re.search(
        r"(\d+(?:\.\d+)?)\s*(MW|KW|W)\b",
        description,
        flags=re.IGNORECASE,
    )

    if not power_match:
        return None

    value = float(
        power_match.group(1)
    )

    unit = power_match.group(2).upper()

    if unit == "MW":
        return value / 1000

    if unit == "KW":
        return value * 1000

    return value


def get_voltages(
    description
):

    # --------------------------------------------------------
    # Plain volt figures, in the order EVVO writes them:
    #     "Diode TVS Single Bi-Dir 188V 600W 2-Pin SMB"
    #     "(TVS/ESD) 24V 45V 300W 27V ESD"
    # The first is the standoff, the second - when there is
    # one - the clamping voltage. kV (ESD strike level) and mV
    # are excluded by the unit guard.
    # --------------------------------------------------------

    voltages = []

    for match in re.finditer(
        r"(?<![A-Za-z0-9.])(\d+(?:\.\d+)?)\s*V(?![A-Za-z0-9])",
        description,
    ):

        try:
            voltages.append(
                float(
                    match.group(1)
                )
            )

        except (ValueError, TypeError):
            continue

    return voltages


def get_voltage_from_part_number(
    part_number
):

    # --------------------------------------------------------
    # EVVO codes the voltage into the name, either in V
    # notation (ZMM5V6 = 5.6 V, PESD3V3V1BCSF = 3.3 V) or
    # after a C (BZX84C10 = 10 V, ZMM55C16 = 16 V).
    # --------------------------------------------------------

    name = str(part_number).strip().upper()

    v_notation = re.search(
        r"(?<!\d)(\d{1,3})V(\d)(?!\d)",
        name,
    )

    if v_notation:

        try:
            return float(
                f"{v_notation.group(1)}."
                f"{v_notation.group(2)}"
            )

        except (ValueError, TypeError):
            return None

    c_notation = re.search(
        r"C(\d{1,3}(?:\.\d+)?)[A-Z]?$",
        name,
    )

    if c_notation:

        try:
            return float(
                c_notation.group(1)
            )

        except (ValueError, TypeError):
            return None

    return None


def get_grade(
    description
):

    text = str(description).lower()

    if "aec" in text or "automotive" in text:
        return "Automotive"

    return "Non-Automotive"


def get_direction(
    description
):

    text = str(description).upper().replace("-", " ")

    if "BI DIR" in text or "BIDIRECTIONAL" in text:
        return "Bidirectional"

    if "UNI DIR" in text or "UNIDIRECTIONAL" in text:
        return "Unidirectional"

    return "-"


def get_clamping_voltage(
    description
):

    # --------------------------------------------------------
    # EVVO marks the clamping voltage with a trailing c:
    #     "... TVS Uni-Dir 5.2V 20Vc Automotive 3-Pin SOT-23"
    # 5.2 V is the standoff, 20 V the clamp. The c keeps this
    # figure out of the plain-volt list in get_voltages.
    # --------------------------------------------------------

    clamp_match = re.search(
        r"(\d+(?:\.\d+)?)\s*VC\b",
        str(description),
        flags=re.IGNORECASE,
    )

    if not clamp_match:
        return None

    try:
        return float(
            clamp_match.group(1)
        )

    except (ValueError, TypeError):
        return None


def get_channels(
    description
):

    # --------------------------------------------------------
    # A stated channel count is used as written. Everything
    # else is a single channel unless the DigiKey export says
    # otherwise - see fetch_evvo_channels in applesauce.
    # --------------------------------------------------------

    channel_match = re.search(
        r"(\d+)\s*-?\s*CHANNEL",
        str(description),
        flags=re.IGNORECASE,
    )

    if channel_match:
        return channel_match.group(1)

    return "1"


# ============================================================
# PARSE TVS ROW
# ============================================================

def parse_evvo_tvs_row(
    part_row,
    source_name,
    part_number,
):

    description = get_text(
        part_row,
        DESCRIPTION_COLUMN,
    )

    encapsulation = get_text(
        part_row,
        ENCAPSULATION_COLUMN,
    )

    specs_result = {
        "Device Name": str(part_number).upper(),
        "Source File": f"EVVO {source_name}",
        "Grade": get_grade(description),
        "Direction": get_direction(description),
        "Channels": get_channels(description),
        "Package": encapsulation,
        "Voltage - Reverse Standoff (Typ)": "-",
        "Voltage - Clamping (Max) @ Ipp": "-",
        "IEC 61000-4-2": "-",
        "IEC 61000-4-5": "-",
        "Capacitance": "-",
        "Power Dissipation (Pd)": "-",
        "Price ($/ku)": "-",
    }

    voltages = get_voltages(
        description
    )

    if voltages:
        specs_result["Voltage - Reverse Standoff (Typ)"] = f"{voltages[0]:g} V"

    else:

        # Nothing in the description, so fall back to the
        # voltage EVVO codes into the name itself.

        name_voltage = get_voltage_from_part_number(
            part_number
        )

        if name_voltage is not None:
            specs_result["Voltage - Reverse Standoff (Typ)"] = f"{name_voltage:g} V"

    # An explicit "20Vc" wins; otherwise the second plain volt
    # figure is the clamp, as in "(TVS/ESD) 24V 45V 300W 27V".

    clamping = get_clamping_voltage(
        description
    )

    if clamping is not None:
        specs_result["Voltage - Clamping (Max) @ Ipp"] = f"{clamping:g} V"

    elif len(voltages) > 1:
        specs_result["Voltage - Clamping (Max) @ Ipp"] = f"{voltages[1]:g} V"

    esd_match = re.search(
        r"(\d+(?:\.\d+)?)\s*KV",
        description,
        flags=re.IGNORECASE,
    )

    if esd_match:
        specs_result["IEC 61000-4-2"] = f"{float(esd_match.group(1)):g} kV"

    surge_match = re.search(
        r"(\d+(?:\.\d+)?)\s*A\b",
        description,
        flags=re.IGNORECASE,
    )

    if surge_match:
        specs_result["IEC 61000-4-5"] = f"{float(surge_match.group(1)):g} A"

    capacitance_match = re.search(
        r"(\d+(?:\.\d+)?)\s*PF",
        description,
        flags=re.IGNORECASE,
    )

    if capacitance_match:
        specs_result["Capacitance"] = f"{float(capacitance_match.group(1)):.2f} pF"

    power = get_power(
        description
    )

    if power is not None:
        specs_result["Power Dissipation (Pd)"] = f"{power:g} W"

    return specs_result


# ============================================================
# PARSE ZENER ROW
# ============================================================

def parse_evvo_zener_row(
    part_row,
    source_name,
    part_number,
):

    description = get_text(
        part_row,
        DESCRIPTION_COLUMN,
    )

    specs_result = {
        "Device Name": str(part_number).upper(),
        "Source File": f"EVVO {source_name}",
        "Grade": get_grade(description),
        "Voltage - Reverse Standoff (Typ)": "-",  # holds Vz for comparison
        "Tolerance": "-",
        "Power Dissipation (Pd)": "-",
        "Package": get_text(part_row, ENCAPSULATION_COLUMN),
        "Price ($/ku)": "-",
    }

    # --------------------------------------------------------
    # Vz is written three ways - "DIODE ZENER 9.1V ...",
    # "VZnom= 16V, IZK= 1mA, LL-34", and "DIODE ZENER 5V6
    # LL34" - and is sometimes only in the name, as with
    # "Ptot= 350mW Vf= 0.9V" on BZX84C3V3A. Vf is a forward
    # drop, not Vz, so a description carrying only that one
    # has to defer to the name.
    # --------------------------------------------------------

    zener_voltage = None

    vz_match = re.search(
        r"VZ\s*(?:NOM)?\s*=?\s*(\d+(?:\.\d+)?)\s*V",
        description,
        flags=re.IGNORECASE,
    )

    if vz_match:
        zener_voltage = float(
            vz_match.group(1)
        )

    if zener_voltage is None:

        v_notation = re.search(
            r"(?<![A-Za-z0-9.])(\d{1,3})V(\d)(?![A-Za-z0-9])",
            description,
        )

        if v_notation:
            zener_voltage = float(
                f"{v_notation.group(1)}."
                f"{v_notation.group(2)}"
            )

    if zener_voltage is None:

        zener_match = re.search(
            r"ZENER\s+(\d+(?:\.\d+)?)\s*V",
            description,
            flags=re.IGNORECASE,
        )

        if zener_match:
            zener_voltage = float(
                zener_match.group(1)
            )

    if zener_voltage is None:
        zener_voltage = get_voltage_from_part_number(
            part_number
        )

    if zener_voltage is not None:
        specs_result["Voltage - Reverse Standoff (Typ)"] = f"{zener_voltage:g} V"

    tolerance_match = re.search(
        r"(\d+(?:\.\d+)?)\s*%",
        description,
    )

    if tolerance_match:
        specs_result["Tolerance"] = f"\u00b1{float(tolerance_match.group(1)):g}%"

    power = get_power(
        description
    )

    if power is not None:
        specs_result["Power Dissipation (Pd)"] = f"{power:g} W"

    return specs_result


# ============================================================
# FETCH SPECS FROM EXCEL
#
# evvo_dfs is {"TVS": df, "Zener": df}.
# ============================================================

def fetch_evvo_specs_from_excel(
    part_number,
    evvo_dfs,
):

    if not part_number or not evvo_dfs:
        return None

    target = str(part_number).strip().lower()

    target_loose = re.sub(
        r"[\W_]+",
        "",
        target,
    )

    search_priority = [
        "TVS",
        "Zener",
    ]

    ordered_keys = [
        key
        for key in search_priority
        if key in evvo_dfs
    ] + [
        key
        for key in evvo_dfs
        if key not in search_priority
    ]

    for df_name in ordered_keys:

        df = evvo_dfs.get(df_name)

        if df is None or df.empty:
            continue

        df.columns = df.columns.str.strip()

        part_column = next(
            (
                column
                for column in df.columns
                if str(column).strip().lower() == MODEL_NUMBER_COLUMN.lower()
            ),
            None,
        )

        if not part_column:
            continue

        column_values = df[part_column].astype(str).str.strip()

        part_rows = df[column_values.str.lower() == target]

        if part_rows.empty:

            loose_values = column_values.str.lower().str.replace(
                r"[\W_]+",
                "",
                regex=True,
            )

            part_rows = df[loose_values == target_loose]

        if part_rows.empty:
            continue

        part_row = part_rows.iloc[0]

        exact_name = safe_strip(
            part_row[part_column]
        )

        # Schottky rectifiers sit in the same tables. TI has no
        # equivalent, so they are refused rather than crossed.

        description = safe_strip(
            get_cell(
                part_row,
                DESCRIPTION_COLUMN,
            )
        )

        if "schottky" in description.lower():

            print(
                f"'{exact_name}' is a Schottky part "
                f"- not crossed."
            )

            return None

        print(
            f"Found specs for Part '{part_number}' "
            f"in EVVO database '{df_name}'."
        )

        if "zener" in df_name.lower():

            return parse_evvo_zener_row(
                part_row,
                df_name,
                exact_name,
            )

        return parse_evvo_tvs_row(
            part_row,
            df_name,
            exact_name,
        )

    print(
        f"Part '{part_number}' not found "
        f"in any EVVO database."
    )

    return None


# ============================================================
# HAS PART
#
# Whether EVVO lists the part at all, so a part refused above
# stops there instead of falling through to DigiKey.
# ============================================================

def evvo_has_part(
    part_number,
    evvo_dfs,
):

    if not part_number or not evvo_dfs:
        return False

    target_loose = re.sub(
        r"[\W_]+",
        "",
        str(part_number).strip().lower(),
    )

    if not target_loose:
        return False

    for df in evvo_dfs.values():

        if df is None or df.empty:
            continue

        df.columns = df.columns.str.strip()

        part_column = next(
            (
                column
                for column in df.columns
                if str(column).strip().lower() == MODEL_NUMBER_COLUMN.lower()
            ),
            None,
        )

        if not part_column:
            continue

        loose_values = (
            df[part_column]
            .astype(str)
            .str.strip()
            .str.lower()
            .str.replace(
                r"[\W_]+",
                "",
                regex=True,
            )
        )

        if (loose_values == target_loose).any():
            return True

    return False


if __name__ == "__main__":
    import traceback

    try:
        evvo_scrape()

    except Exception:
        print(
            "\n"
            + "=" * 70
        )

        print(
            "SCRAPER ERROR"
        )

        print(
            "=" * 70
        )

        traceback.print_exc()

    finally:
        print()

        print(
            "Window will stay open "
            "so you can read the output."
        )

        input(
            "Press ENTER to close..."
        )

if __name__ == "__main__":
    import traceback

    try:
        evvo_scrape()

    except Exception:
        print("\n" + "=" * 70)
        print("SCRAPER ERROR")
        print("=" * 70)
        traceback.print_exc()

    finally:
        print("\nWindow will stay open so you can read the output.")
        input("Press ENTER to close...")